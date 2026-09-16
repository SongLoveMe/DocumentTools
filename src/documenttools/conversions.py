from __future__ import annotations

import tempfile
from pathlib import Path

from .engines import (
    AppKind,
    ConversionEngineError,
    EngineKind,
    LocalEngineInfo,
    app_kind_for_path,
    convert_with_local_engine,
    engine_for_app,
)
from .paths import classify_file, unique_path

# Guidance shown when no local Office/WPS engine is available.  The user is
# intentionally pointed at online services instead of a built-in fallback:
# DocumentTools never ships a renderer of its own.
ONLINE_CONVERSION_SERVICES = (
    ("Smallpdf", "https://smallpdf.com/cn"),
    ("iLovePDF", "https://www.ilovepdf.com/zh-cn"),
    ("PDF24 Tools", "https://tools.pdf24.org/zh/all-tools"),
)


class ConversionError(RuntimeError):
    pass


def image_to_pdf(sources: list[str | Path], output: str | Path) -> None:
    try:
        from PIL import Image  # type: ignore[import-not-found]
    except ImportError as exc:
        raise ConversionError("Pillow is not installed in the DocumentTools environment.") from exc
    paths = [Path(source) for source in sources]
    if not paths:
        raise ConversionError("Add at least one image.")
    images = []
    try:
        for path in paths:
            if classify_file(path) != "image":
                raise ConversionError(f"Unsupported image: {path.name}")
            with Image.open(path) as source_image:
                for frame in range(getattr(source_image, "n_frames", 1)):
                    source_image.seek(frame)
                    image = source_image.convert("RGB")
                    images.append(image.copy())
        if not images:
            raise ConversionError("No image pages were available.")
        images[0].save(str(output), "PDF", save_all=True, append_images=images[1:], resolution=150.0)
    finally:
        for image in images:
            image.close()


def convert_file_to_pdf(
    source: str | Path,
    output_directory: str | Path,
    output_stem: str | None = None,
    *,
    engines: list[LocalEngineInfo] | tuple[LocalEngineInfo, ...] | None = None,
    engine_kind: EngineKind | None = None,
    timeout_seconds: int = 120,
) -> tuple[Path, str]:
    """Convert a single file to PDF and report the component that did it.

    Images are handled by Pillow.  Office documents require a locally installed
    Microsoft Office or WPS engine, which is driven through COM automation;
    when none is available :class:`ConversionError` is raised so the UI can
    offer online conversion instead.
    """
    source_path = Path(source)
    output = unique_path(output_directory, output_stem or source_path.stem, ".pdf")
    file_kind = classify_file(source_path)
    if file_kind == "image":
        image_to_pdf([source_path], output)
        return output, "Pillow"
    if file_kind == "office":
        return _convert_office(
            source_path, output, engines, engine_kind, timeout_seconds
        )
    raise ConversionError(f"不支持的输入格式：{source_path.name}")


def _convert_office(
    source_path: Path,
    output: Path,
    engines: list[LocalEngineInfo] | tuple[LocalEngineInfo, ...] | None,
    engine_kind: EngineKind | None,
    timeout_seconds: int,
) -> tuple[Path, str]:
    from .engines import detect_local_engines

    app_kind: AppKind | None = app_kind_for_path(source_path)
    if app_kind is None:
        raise ConversionError(f"不支持的 Office 格式：{source_path.suffix}")

    detected = tuple(engines) if engines is not None else detect_local_engines()
    if not detected:
        raise ConversionError("未检测到本机 Microsoft Office 或 WPS，无法转换该文件。")

    candidates: list[LocalEngineInfo] = []
    if engine_kind is not None:
        # An explicitly chosen engine is never silently swapped for the other.
        chosen = engine_for_app(detected, app_kind, preferred=engine_kind)
        if chosen is not None and chosen.kind is engine_kind:
            candidates.append(chosen)
    else:
        for kind in (EngineKind.OFFICE, EngineKind.WPS):
            candidate = engine_for_app(detected, app_kind, preferred=kind)
            if candidate is not None and candidate not in candidates:
                candidates.append(candidate)

    if not candidates:
        raise ConversionError(
            f"未检测到可转换 {source_path.suffix} 的 Microsoft Office 或 WPS 组件。"
        )

    failures: list[str] = []
    for candidate in candidates:
        binding = candidate.binding_for(app_kind)
        try:
            convert_with_local_engine(
                source_path,
                output,
                engine_kind=candidate.kind,
                app_kind=app_kind,
                prog_id=binding.prog_id if binding else None,
                timeout_seconds=timeout_seconds,
            )
        except ConversionEngineError as exc:
            failures.append(f"{candidate.label()}：{exc}")
            continue
        return output, f"{candidate.display_name} {candidate.version}".strip()

    if output.exists():
        output.unlink()
    raise ConversionError("；".join(failures))


def convert_pdf_to_word(
    source: str | Path, output_directory: str | Path, output_stem: str | None = None
) -> Path:
    """Rebuild a simplified Word document from extracted PDF text and images.

    This is intentionally a reconstruction rather than a faithful conversion:
    PDF stores positioned glyphs, not editable structure, so paragraphs, tables
    and columns cannot be recovered reliably.  Output keeps the readable text
    and embedded images so the result stays editable in Word.
    """
    try:
        import fitz  # type: ignore[import-not-found]
        from docx import Document  # type: ignore[import-not-found]
        from docx.shared import Inches  # type: ignore[import-not-found]
    except ImportError as exc:
        raise ConversionError("PDF 转 Word 所需的本地依赖未安装，请安装 requirements.lock。") from exc
    source_path = Path(source)
    if classify_file(source_path) != "pdf":
        raise ConversionError("PDF 转 Word 只接受 PDF 文件。")
    output = unique_path(output_directory, output_stem or source_path.stem, ".docx")

    document = fitz.open(source_path)
    try:
        if document.needs_pass and not document.authenticate(""):
            raise ConversionError(f"{source_path.name} 需要打开密码，无法猜测密码。")
        if document.page_count == 0:
            raise ConversionError(f"{source_path.name} 没有页面。")

        word = Document()
        for page_index, page in enumerate(document, start=1):
            if page_index > 1:
                word.add_page_break()
            text = page.get_text("text").strip()
            if text:
                for line in text.splitlines():
                    word.add_paragraph(line)
            for image in page.get_images(full=True):
                try:
                    extracted = document.extract_image(image[0])
                    word.add_picture(io.BytesIO(extracted["image"]), width=Inches(5.5))
                except Exception:
                    continue
        try:
            word.save(str(output))
        except Exception as exc:
            raise ConversionError(f"无法写入 Word 文件：{exc}") from exc
    finally:
        document.close()
    return output


def convert_pdf_to_ppt(source: str | Path, output_directory: str | Path, output_stem: str | None = None) -> Path:
    try:
        import fitz  # type: ignore[import-not-found]
        from pptx import Presentation  # type: ignore[import-not-found]
        from pptx.util import Inches  # type: ignore[import-not-found]
    except ImportError as exc:
        raise ConversionError("PDF to PowerPoint support is not installed. Run pip install -r requirements.lock.") from exc
    source_path = Path(source)
    if classify_file(source_path) != "pdf":
        raise ConversionError("PDF to PowerPoint accepts only PDF files.")
    output = unique_path(output_directory, output_stem or source_path.stem, ".pptx")
    try:
        document = fitz.open(source_path)
    except Exception as exc:
        raise ConversionError(f"Could not open {source_path.name}: {exc}") from exc
    if document.needs_pass and not document.authenticate(""):
        document.close()
        raise ConversionError(f"{source_path.name} 需要打开密码，无法猜测密码。")
    if document.page_count == 0:
        document.close()
        raise ConversionError(f"{source_path.name} has no pages.")
    try:
        first = document[0].rect
        presentation = Presentation()
        presentation.slide_width = Inches(13.333)
        presentation.slide_height = int(presentation.slide_width * first.height / first.width)
        blank = presentation.slide_layouts[6]
        with tempfile.TemporaryDirectory(prefix="documenttools-ppt-") as temporary:
            for index, page in enumerate(document):
                raster = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
                image_path = Path(temporary) / f"page-{index + 1}.png"
                raster.save(str(image_path))
                slide = presentation.slides.add_slide(blank)
                page_ratio = page.rect.width / page.rect.height
                slide_ratio = presentation.slide_width / presentation.slide_height
                if page_ratio >= slide_ratio:
                    width = presentation.slide_width
                    height = int(width / page_ratio)
                    left, top = 0, int((presentation.slide_height - height) / 2)
                else:
                    height = presentation.slide_height
                    width = int(height * page_ratio)
                    left, top = int((presentation.slide_width - width) / 2), 0
                slide.shapes.add_picture(str(image_path), left, top, width=width, height=height)
        presentation.save(str(output))
    finally:
        document.close()
    return output


def _open_pdf_document(source: str | Path):
    try:
        import fitz  # type: ignore[import-not-found]
    except ImportError as exc:
        raise ConversionError("PDF 转换支持未安装，请使用 documenttools 环境安装依赖。") from exc
    source_path = Path(source)
    if classify_file(source_path) != "pdf" or not source_path.is_file():
        raise ConversionError("请选择可读取的 PDF 文件。")
    try:
        document = fitz.open(source_path)
    except Exception as exc:
        raise ConversionError(f"无法打开 {source_path.name}：{exc}") from exc
    if document.needs_pass and not document.authenticate(""):
        document.close()
        raise ConversionError(f"{source_path.name} 需要打开密码，无法猜测密码。")
    if document.page_count == 0:
        document.close()
        raise ConversionError(f"{source_path.name} 没有页面。")
    return source_path, document


def _unique_directory(parent: Path, stem: str) -> Path:
    candidate = parent / stem
    sequence = 1
    while candidate.exists():
        candidate = parent / f"{stem} ({sequence})"
        sequence += 1
    candidate.mkdir(parents=True, exist_ok=False)
    return candidate


def convert_pdf_to_excel(source: str | Path, output_directory: str | Path | None = None, output_stem: str | None = None) -> Path:
    """Extract detectable digital-PDF tables into one worksheet per table."""
    try:
        import xlsxwriter  # type: ignore[import-not-found]
    except ImportError as exc:
        raise ConversionError("PDF 转 Excel 支持未安装，请安装 requirements.lock 中的 XlsxWriter。") from exc
    source_path, document = _open_pdf_document(source)
    directory = Path(output_directory) if output_directory else source_path.parent
    output = unique_path(directory, output_stem or f"{source_path.stem}_表格", ".xlsx")
    extracted: list[tuple[int, list[list[object]]]] = []
    try:
        for page_index, page in enumerate(document, start=1):
            try:
                tables = page.find_tables()
            except Exception as exc:
                raise ConversionError(f"第 {page_index} 页表格识别失败：{exc}") from exc
            for table in tables.tables:
                values = table.extract()
                if values and any(any(cell not in (None, "") for cell in row) for row in values):
                    extracted.append((page_index, values))
        if not extracted:
            raise ConversionError("未识别到可提取的表格。仅支持带文本层的规则表格，不支持扫描件或复杂视觉排版。")
        directory.mkdir(parents=True, exist_ok=True)
        workbook = xlsxwriter.Workbook(str(output))
        try:
            for number, (page_index, rows) in enumerate(extracted, start=1):
                worksheet = workbook.add_worksheet(f"第 {page_index} 页-表 {number}")
                for row_index, row in enumerate(rows):
                    for column_index, value in enumerate(row):
                        worksheet.write(row_index, column_index, "" if value is None else str(value))
                worksheet.freeze_panes(1, 0)
                for column_index, row in enumerate(rows[0] if rows else []):
                    values = ["" if source_row[column_index] is None else str(source_row[column_index]) for source_row in rows if column_index < len(source_row)]
                    worksheet.set_column(column_index, column_index, min(max(max((len(value) for value in values), default=8) + 2, 10), 40))
        finally:
            workbook.close()
    except Exception:
        if output.exists():
            output.unlink()
        raise
    finally:
        document.close()
    return output


def convert_pdf_to_images(
    source: str | Path,
    *,
    image_format: str = "png",
    dpi: int = 150,
    output_directory: str | Path | None = None,
) -> list[Path]:
    """Render each PDF page to PNG or high-quality JPEG in a unique folder."""
    import fitz  # type: ignore[import-not-found]
    if dpi not in {72, 150, 300}:
        raise ConversionError("图片分辨率仅支持 72、150 或 300 DPI。")
    normalized_format = image_format.lower().replace("jpeg", "jpg")
    if normalized_format not in {"png", "jpg"}:
        raise ConversionError("图片格式仅支持 PNG 或 JPG。")
    source_path, document = _open_pdf_document(source)
    if output_directory:
        target = _unique_directory(Path(output_directory), f"{source_path.stem}_{normalized_format.upper()}")
    else:
        target = _unique_directory(source_path.parent, f"{source_path.stem}_{normalized_format.upper()}")
    scale = dpi / 72
    results: list[Path] = []
    try:
        for page_index, page in enumerate(document, start=1):
            target_path = target / f"{source_path.stem}_{page_index:04d}.{normalized_format}"
            pixmap = page.get_pixmap(matrix=fitz.Matrix(scale, scale), alpha=False)
            if normalized_format == "jpg":
                pixmap.save(str(target_path), jpg_quality=95)
            else:
                pixmap.save(str(target_path))
            results.append(target_path)
    finally:
        document.close()
    return results
